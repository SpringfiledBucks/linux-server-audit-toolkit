"""Fixed-command, read-only Linux and BMC collection."""
from __future__ import annotations

import argparse
import csv
import ipaddress
import json
import math
import os
import re
import socket
import subprocess
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


COMMANDS = {
    "hostname": "hostname",
    "uptime": "uptime -p",
    "kernel": "uname -r",
    "os_release": "cat /etc/os-release",
    "cpu": "lscpu",
    "memory": "free -b",
    "filesystem_usage": "df -P",
    "block_devices": "lsblk -o NAME,SIZE,TYPE,MOUNTPOINT -P",
    "ip_addresses": "ip -o addr show",
    "default_route": "ip route show default",
    "failed_services": "systemctl --failed --no-legend --plain",
}
COMMANDS = {field: "LC_ALL=C " + command for field, command in COMMANDS.items()}
LINUX_FIELDS = tuple(COMMANDS)
BMC_FIELDS = ("fru_device", "manufacturer", "product", "serial_number")
ERRORS = ("OK", "SSH_AUTH_FAILED", "SSH_TIMEOUT", "BMC_AUTH_FAILED", "BMC_TIMEOUT", "COMMAND_FAILED", "PARSE_FAILED")
HOST_RE = re.compile(r"^[A-Za-z0-9](?:[A-Za-z0-9.-]{0,251}[A-Za-z0-9])?$")
NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}$")
MAX_OUTPUT = 16_000
ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class AuditError(Exception):
    def __init__(self, code: str, field: str = "", detail: str = "") -> None:
        if code not in ERRORS:
            raise ValueError("invalid error code")
        self.code, self.field, self.detail = code, field, detail
        super().__init__(code)


def is_timeout_exception(error: BaseException) -> bool:
    seen = set()
    pending = [error]
    while pending:
        current = pending.pop()
        if id(current) in seen:
            continue
        seen.add(id(current))
        if isinstance(current, (TimeoutError, socket.timeout)):
            return True
        pending.extend(item for item in (current.__cause__, current.__context__) if item is not None)
    return False


@dataclass(frozen=True)
class Host:
    name: str
    host: str | None
    ssh_user: str | None
    bmc_host: str | None


def valid_host(value: object) -> bool:
    if not isinstance(value, str) or not HOST_RE.fullmatch(value) or ".." in value:
        return False
    try:
        ipaddress.ip_address(value)
        return True
    except ValueError:
        if re.fullmatch(r"[0-9.]+", value):
            return False
        return all(label and len(label) <= 63 and not label.startswith("-") and not label.endswith("-") for label in value.split("."))


def read_inventory(path: Path) -> list[Host]:
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict) or set(raw) != {"hosts"} or not isinstance(raw["hosts"], list) or not raw["hosts"]:
        raise ValueError("inventory needs a nonempty hosts list")
    hosts: list[Host] = []
    names: set[str] = set()
    for item in raw["hosts"]:
        if not isinstance(item, dict) or set(item) - {"name", "host", "ssh_user", "bmc_host"}:
            raise ValueError("invalid inventory host fields")
        name, host, user, bmc = (item.get(k) for k in ("name", "host", "ssh_user", "bmc_host"))
        if not isinstance(name, str) or not NAME_RE.fullmatch(name) or name in names:
            raise ValueError("invalid or duplicate host name")
        if not host and not bmc:
            raise ValueError("host must have Linux SSH and/or BMC target")
        if host and (not valid_host(host) or not isinstance(user, str) or not NAME_RE.fullmatch(user)):
            raise ValueError("invalid SSH target or user")
        if bmc and not valid_host(bmc):
            raise ValueError("invalid BMC target")
        if not host and user:
            raise ValueError("ssh_user needs host")
        names.add(name)
        hosts.append(Host(name, host, user, bmc))
    return hosts


class Provider(Protocol):
    def linux(self, host: Host, field: str) -> str: ...
    def bmc(self, host: Host) -> str: ...


class MockProvider:
    """Public synthetic responses only; fixtures are bundled with the demo."""

    def __init__(self, path: Path) -> None:
        self.data = json.loads(path.read_text(encoding="utf-8"))

    def linux(self, host: Host, field: str) -> str:
        entry = self.data.get(host.name, {}).get("linux", {})
        if entry.get("error"):
            raise AuditError(entry["error"], field)
        item = entry.get(field)
        if isinstance(item, dict):
            raise AuditError(item.get("error", "COMMAND_FAILED"), field)
        if item is None:
            raise AuditError("COMMAND_FAILED", field)
        return str(item)

    def bmc(self, host: Host) -> str:
        entry = self.data.get(host.name, {}).get("bmc", {})
        if entry.get("error"):
            raise AuditError(entry["error"], "fru")
        return str(entry.get("fru", ""))


class RealProvider:
    """Strict SSH host keys and a single fixed ipmitool FRU command."""

    def __init__(self, timeout: int, known_hosts: Path | None = None) -> None:
        self.timeout = timeout
        self.known_hosts = known_hosts

    def linux(self, host: Host, field: str) -> str:
        import paramiko

        client = paramiko.SSHClient()
        connected = False
        try:
            if self.known_hosts:
                client.load_host_keys(str(self.known_hosts))
            else:
                client.load_system_host_keys()
            client.set_missing_host_key_policy(paramiko.RejectPolicy())
            kwargs = {"hostname": host.host, "username": host.ssh_user, "timeout": self.timeout,
                      "auth_timeout": self.timeout, "banner_timeout": self.timeout,
                      "look_for_keys": False, "allow_agent": False}
            key_path = os.getenv("AUDIT_SSH_KEY_PATH")
            password = os.getenv("AUDIT_SSH_PASSWORD")
            if key_path:
                kwargs["key_filename"] = key_path
            elif password:
                kwargs["password"] = password
            client.connect(**kwargs)
            connected = True
            _, stdout, stderr = client.exec_command(COMMANDS[field], timeout=self.timeout, get_pty=False)
            stdout.channel.settimeout(self.timeout)
            output = stdout.read(MAX_OUTPUT + 1)
            if len(output) > MAX_OUTPUT:
                raise AuditError("COMMAND_FAILED", field)
            error_output = stderr.read(MAX_OUTPUT + 1)
            if len(error_output) > MAX_OUTPUT:
                raise AuditError("COMMAND_FAILED", field)
            status = stdout.channel.recv_exit_status()
            if status != 0:
                raise AuditError("COMMAND_FAILED", field)
            return output.decode("utf-8", errors="replace")
        except paramiko.AuthenticationException:
            raise AuditError("SSH_AUTH_FAILED", "connection") from None
        except (TimeoutError, OSError) as exc:
            if is_timeout_exception(exc):
                raise AuditError("SSH_TIMEOUT", "connection" if not connected else field) from None
            raise AuditError("COMMAND_FAILED", "connection" if not connected else field) from None
        except AuditError:
            raise
        except Exception as exc:
            # Never include provider exception text: it can contain an endpoint or secret.
            code = "SSH_TIMEOUT" if is_timeout_exception(exc) else "COMMAND_FAILED"
            raise AuditError(code, "connection" if not connected else field) from None
        finally:
            client.close()

    def bmc(self, host: Host) -> str:
        users = [x.strip() for x in os.getenv("AUDIT_BMC_USERS", "").split(",") if x.strip()]
        if not users:
            raise AuditError("BMC_AUTH_FAILED", "fru", "no credential configured")
        attempted = False
        for index, user in enumerate(users, start=1):
            if not NAME_RE.fullmatch(user):
                continue
            password = os.getenv(f"AUDIT_BMC_PASSWORD_{index}")
            if not password:
                continue
            attempted = True
            env = {key: value for key, value in os.environ.items()
                   if not key.startswith("AUDIT_BMC_PASSWORD_") and key not in ("AUDIT_SSH_PASSWORD", "IPMI_PASSWORD")}
            env["IPMI_PASSWORD"] = password
            try:
                result = subprocess.run(
                    ["ipmitool", "-I", "lanplus", "-H", host.bmc_host, "-U", user, "-E", "fru"],
                    capture_output=True, text=True, timeout=self.timeout, env=env, check=False,
                )
            except subprocess.TimeoutExpired:
                raise AuditError("BMC_TIMEOUT", "fru") from None
            except OSError:
                raise AuditError("COMMAND_FAILED", "fru") from None
            finally:
                env.pop("IPMI_PASSWORD", None)
            if result.returncode == 0:
                return result.stdout[:MAX_OUTPUT]
            error = result.stderr.lower()
            if any(term in error for term in ("password", "authentication", "unauthorized", "rakp 2", "invalid user")):
                continue
            if any(term in error for term in ("timed out", "timeout", "unreachable")):
                raise AuditError("BMC_TIMEOUT", "fru")
            raise AuditError("COMMAND_FAILED", "fru")
        raise AuditError("BMC_AUTH_FAILED", "fru", "credentials rejected" if attempted else "no credential configured")


def parse_linux(field: str, raw: str) -> str:
    raw = raw.strip()[:MAX_OUTPUT]
    if field == "os_release":
        match = re.search(r"^PRETTY_NAME=(.+)$", raw, re.M)
        raw = match.group(1).strip().strip('"\'') if match else ""
    elif field == "cpu":
        match = re.search(r"^(?:Model name|Architecture):\s*(.+)$", raw, re.M)
        raw = match.group(1) if match else ""
    elif field == "memory":
        match = re.search(r"^Mem:\s+(\d+)", raw, re.M)
        raw = f"{int(match.group(1)) / (1024**3):.1f} GiB total" if match else ""
    elif field == "failed_services" and not raw:
        return "none"
    if not raw:
        raise AuditError("PARSE_FAILED", field)
    return raw


def parse_fru(raw: str) -> dict[str, str]:
    headers = list(re.finditer(r"^[ \t]*FRU Device Description[ \t]*:[ \t]*([^\r\n]+)$", raw, re.M | re.I))
    if len(headers) > 1:
        selected = [index for index, header in enumerate(headers)
                    if header.group(1).strip().lower() == "builtin fru device (id 0)"]
        if len(selected) != 1:
            raise AuditError("PARSE_FAILED", "fru", "ambiguous FRU device")
        index = selected[0]
    else:
        index = 0
    if headers:
        raw = raw[headers[index].end():headers[index + 1].start() if index + 1 < len(headers) else len(raw)]
        device = headers[index].group(1).strip()
    else:
        device = "unspecified"
    values = {"fru_device": device}
    for key, pattern in {
        "manufacturer": r"^[ \t]*(?:Product|Board) Manufacturer[ \t]*:[ \t]*([^\r\n]*)$",
        "product": r"^[ \t]*Product Name[ \t]*:[ \t]*([^\r\n]*)$",
        "serial_number": r"^[ \t]*(?:Product|Board) Serial[ \t]*:[ \t]*([^\r\n]*)$",
    }.items():
        match = re.search(pattern, raw, re.M | re.I)
        values[key] = match.group(1).strip() if match else ""
    if not any(values[field] for field in ("manufacturer", "product", "serial_number")):
        raise AuditError("PARSE_FAILED", "fru")
    return values


def collect(hosts: list[Host], provider: Provider) -> tuple[list[dict], list[dict], list[dict]]:
    linux_rows, bmc_rows, failures = [], [], []
    for host in hosts:
        linux = {"name": host.name, "host": host.host or ""}
        bmc = {"name": host.name, "bmc_host": host.bmc_host or ""}
        linux_errors, bmc_errors = [], []
        if host.host:
            for field in LINUX_FIELDS:
                try:
                    linux[field] = parse_linux(field, provider.linux(host, field))
                except AuditError as exc:
                    linux[field] = ""
                    linux_errors.append(exc.code)
                    failures.append({"name": host.name, "component": "linux", "field": exc.field or field,
                                     "code": exc.code, "detail": exc.detail})
                    if exc.field == "connection" or exc.code in ("SSH_AUTH_FAILED", "SSH_TIMEOUT"):
                        # Connection-level failure: one bounded attempt per host.
                        for remaining in LINUX_FIELDS[LINUX_FIELDS.index(field) + 1:]:
                            linux[remaining] = ""
                        break
                except Exception:
                    linux[field] = ""
                    linux_errors.append("COMMAND_FAILED")
                    failures.append({"name": host.name, "component": "linux", "field": field,
                                     "code": "COMMAND_FAILED", "detail": "provider error"})
            linux["status"] = "OK" if not linux_errors else ("PARTIAL" if any(linux.get(f) for f in LINUX_FIELDS) else "FAILED")
            linux_rows.append(linux)
        if host.bmc_host:
            try:
                bmc.update(parse_fru(provider.bmc(host)))
            except AuditError as exc:
                bmc_errors.append(exc.code)
                failures.append({"name": host.name, "component": "bmc", "field": exc.field or "fru",
                                 "code": exc.code, "detail": exc.detail})
            except Exception:
                bmc_errors.append("COMMAND_FAILED")
                failures.append({"name": host.name, "component": "bmc", "field": "fru",
                                 "code": "COMMAND_FAILED", "detail": "provider error"})
            for field in BMC_FIELDS:
                bmc.setdefault(field, "")
                if not bmc[field] and not any(x in bmc_errors for x in ("BMC_AUTH_FAILED", "BMC_TIMEOUT", "COMMAND_FAILED")):
                    bmc_errors.append("PARSE_FAILED")
                    failures.append({"name": host.name, "component": "bmc", "field": field,
                                     "code": "PARSE_FAILED", "detail": "missing FRU field"})
            bmc["status"] = "OK" if not bmc_errors else ("PARTIAL" if any(bmc[f] for f in BMC_FIELDS) else "FAILED")
            bmc_rows.append(bmc)
    return linux_rows, bmc_rows, failures


def safe_cell(value: object) -> str:
    text = ILLEGAL_XLSX.sub("", str(value if value is not None else ""))
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def write_csv(path: Path, fields: tuple[str, ...], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: safe_cell(row.get(k, "")) for k in fields})


def write_reports(output: Path, hosts: list[Host], linux: list[dict], bmc: list[dict], failures: list[dict]) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    linux_fields = ("name", "host", "status", *LINUX_FIELDS)
    bmc_fields = ("name", "bmc_host", "status", *BMC_FIELDS)
    failure_fields = ("name", "component", "field", "code", "detail")
    write_csv(output / "linux_status.csv", linux_fields, linux)
    write_csv(output / "bmc_inventory.csv", bmc_fields, bmc)
    write_csv(output / "failures.csv", failure_fields, failures)
    counts = Counter(row["code"] for row in failures)
    summary = {"hosts": len(hosts), "linux_targets": len(linux), "bmc_targets": len(bmc),
               "linux_ok": sum(x["status"] == "OK" for x in linux),
               "bmc_ok": sum(x["status"] == "OK" for x in bmc),
               "failure_count": len(failures), "error_counts": dict(sorted(counts.items())),
               "result": "OK" if not failures else "PARTIAL"}
    (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    book = Workbook()
    book.remove(book.active)
    for title, fields, rows in (
        ("Summary", ("metric", "value"), [{"metric": k, "value": json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v} for k, v in summary.items()]),
        ("Linux Status", linux_fields, linux), ("BMC Inventory", bmc_fields, bmc),
        ("Failures", failure_fields, failures),
    ):
        sheet = book.create_sheet(title)
        sheet.append(list(fields))
        for row in rows:
            sheet.append([value if isinstance((value := row.get(key, "")), (int, float)) else safe_cell(value) for key in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 28
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
        for row in sheet.iter_rows(min_row=2):
            lines = max(sum(max(1, math.ceil(len(part) / max(1, sheet.column_dimensions[cell.column_letter].width - 2)))
                            for part in str(cell.value or "").split("\n")) for cell in row)
            sheet.row_dimensions[row[0].row].height = min(409, max(22, lines * 16))
    book.save(output / "server_inventory.xlsx")
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Read-only Linux/BMC inventory")
    parser.add_argument("--inventory", type=Path, required=True)
    parser.add_argument("--provider", choices=("mock", "real"), default="mock")
    parser.add_argument("--output", type=Path, default=Path("output"))
    parser.add_argument("--mock-data", type=Path, default=Path(__file__).resolve().parent.parent / "demo" / "mock_responses.json")
    parser.add_argument("--timeout", type=int, default=8)
    parser.add_argument("--known-hosts", type=Path)
    args = parser.parse_args(argv)
    if not 1 <= args.timeout <= 120:
        parser.error("timeout must be 1..120 seconds")
    try:
        hosts = read_inventory(args.inventory)
        provider = MockProvider(args.mock_data) if args.provider == "mock" else RealProvider(args.timeout, args.known_hosts)
        linux, bmc, failures = collect(hosts, provider)
        summary = write_reports(args.output, hosts, linux, bmc, failures)
    except (OSError, ValueError, yaml.YAMLError, json.JSONDecodeError):
        parser.exit(2, "Invalid inventory, mock data, or output path.\n")
    print(json.dumps(summary, ensure_ascii=False))
    return 0
