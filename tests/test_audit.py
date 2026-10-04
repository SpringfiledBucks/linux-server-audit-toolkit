import csv
import json
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.audit import (
    AuditError, Host, MockProvider, RealProvider, collect, main, read_inventory,
    parse_fru, safe_cell, valid_host, write_reports,
)


ROOT = Path(__file__).resolve().parents[1]


def test_demo_all_hosts_and_partial_fields(tmp_path):
    hosts = read_inventory(ROOT / "demo" / "hosts.yaml")
    linux, bmc, failures = collect(hosts, MockProvider(ROOT / "demo" / "mock_responses.json"))
    summary = write_reports(tmp_path, hosts, linux, bmc, failures)
    assert len(hosts) == len(linux) == len(bmc) == 4
    assert summary["linux_ok"] == 1
    assert summary["bmc_ok"] == 2
    assert summary["result"] == "PARTIAL"
    assert {x["code"] for x in failures} >= {"SSH_TIMEOUT", "SSH_AUTH_FAILED", "BMC_AUTH_FAILED", "COMMAND_FAILED", "PARSE_FAILED"}
    assert linux[0]["status"] == "OK" and linux[0]["failed_services"] == "none"
    assert linux[1]["status"] == "FAILED" and bmc[1]["status"] == "OK"
    assert linux[3]["status"] == "PARTIAL" and bmc[3]["status"] == "PARTIAL"
    assert bmc[3]["serial_number"] == ""
    assert load_workbook(tmp_path / "server_inventory.xlsx").sheetnames == [
        "Summary", "Linux Status", "BMC Inventory", "Failures"
    ]
    assert all((tmp_path / name).exists() for name in (
        "linux_status.csv", "bmc_inventory.csv", "failures.csv", "summary.json"
    ))


def test_bmc_timeout_and_parse_failure_are_isolated():
    class FixtureProvider:
        def linux(self, host, field):
            return "demo" if field not in ("os_release", "cpu", "memory") else {
                "os_release": 'PRETTY_NAME="Demo OS"', "cpu": "Model name: Demo CPU", "memory": "Mem: 1073741824"
            }[field]

        def bmc(self, host):
            if host.name == "timeout":
                raise AuditError("BMC_TIMEOUT", "fru")
            return "invalid FRU response"

    hosts = [Host("timeout", None, None, "192.0.2.110"),
             Host("parse", "192.0.2.10", "demo", "192.0.2.111")]
    linux, bmc, failures = collect(hosts, FixtureProvider())
    assert len(linux) == 1 and linux[0]["status"] == "OK"
    assert [x["status"] for x in bmc] == ["FAILED", "FAILED"]
    assert {x["code"] for x in failures} == {"BMC_TIMEOUT", "PARSE_FAILED"}


def test_real_bmc_retries_only_auth_and_keeps_password_out_of_args(monkeypatch):
    monkeypatch.setenv("AUDIT_BMC_USERS", "user1,user2")
    monkeypatch.setenv("AUDIT_BMC_PASSWORD_1", "secret-one")
    monkeypatch.setenv("AUDIT_BMC_PASSWORD_2", "secret-two")
    calls = []

    def fake_run(args, **kwargs):
        calls.append((args, kwargs["env"]["IPMI_PASSWORD"]))
        return subprocess.CompletedProcess(args, 1 if len(calls) == 1 else 0,
                                           "Product Name : Demo Server" if len(calls) == 2 else "",
                                           "authentication failed" if len(calls) == 1 else "")

    monkeypatch.setattr(subprocess, "run", fake_run)
    assert "Demo Server" in RealProvider(3).bmc(Host("demo", None, None, "192.0.2.110"))
    assert len(calls) == 2
    assert all(call[0][-1] == "fru" and call[0][-2] == "-E" for call in calls)
    assert all("secret" not in " ".join(call[0]) for call in calls)
    assert [call[1] for call in calls] == ["secret-one", "secret-two"]

    calls.clear()

    def non_auth(args, **kwargs):
        calls.append(args)
        return subprocess.CompletedProcess(args, 1, "", "device unavailable")

    monkeypatch.setattr(subprocess, "run", non_auth)
    with pytest.raises(AuditError) as exc:
        RealProvider(3).bmc(Host("demo", None, None, "192.0.2.110"))
    assert exc.value.code == "COMMAND_FAILED" and len(calls) == 1


def test_inventory_rejects_unsafe_targets(tmp_path):
    for address in ("999.999.999.999", "192.0.2.10;reboot", "bad..name"):
        assert not valid_host(address)
    path = tmp_path / "hosts.yaml"
    path.write_text("hosts:\n  - name: x\n    host: 192.0.2.10;reboot\n    ssh_user: demo\n", encoding="utf-8")
    with pytest.raises(ValueError):
        read_inventory(path)


def test_csv_and_excel_formula_escaping(tmp_path):
    host = Host("demo", "192.0.2.10", "demo", None)
    linux = [{"name": "demo", "host": "192.0.2.10", "status": "PARTIAL", "hostname": "=1+1\x01"}]
    write_reports(tmp_path, [host], linux, [], [])
    with (tmp_path / "linux_status.csv").open(encoding="utf-8-sig", newline="") as stream:
        row = next(csv.DictReader(stream))
    assert row["hostname"] == "'=1+1"
    sheet = load_workbook(tmp_path / "server_inventory.xlsx")["Linux Status"]
    assert sheet["D2"].value == "'=1+1"
    assert safe_cell(" @SUM(1)") == "' @SUM(1)"
    assert safe_cell("normal\x01value") == "normalvalue"
    assert load_workbook(tmp_path / "server_inventory.xlsx")["Summary"]["B2"].data_type == "n"


def test_connection_failure_one_attempt_and_next_host_continues():
    calls = []

    class Provider:
        def linux(self, host, field):
            calls.append((host.name, field))
            if host.name == "bad":
                raise AuditError("SSH_TIMEOUT", field)
            return {"os_release": 'PRETTY_NAME="Demo OS"', "cpu": "Model name: Demo CPU",
                    "memory": "Mem: 1073741824", "failed_services": "none"}.get(field, "demo")

        def bmc(self, host):
            return "  Product Manufacturer : Demo\n  Product Name : Mock\n  Product Serial : DEMO-001"

    hosts = [Host("bad", "192.0.2.10", "demo", "192.0.2.110"),
             Host("good", "192.0.2.11", "demo", None)]
    linux, bmc, failures = collect(hosts, Provider())
    assert sum(name == "bad" for name, _ in calls) == 1
    assert linux[0]["status"] == "FAILED" and linux[1]["status"] == "OK"
    assert bmc[0]["status"] == "OK"
    assert len(failures) == 1 and failures[0]["code"] == "SSH_TIMEOUT"


def test_fru_multiple_devices_selects_builtin_only():
    raw = ("FRU Device Description : PSU (ID 1)\n"
           " Product Manufacturer : PSU Vendor\n Product Name : Power Unit\n Product Serial : PSU-999\n"
           "FRU Device Description : Builtin FRU Device (ID 0)\n"
           " Product Manufacturer : Server Vendor\n Product Name : Server\n Product Serial : SYSTEM-123\n")
    result = parse_fru(raw)
    assert result == {"fru_device": "Builtin FRU Device (ID 0)",
                      "manufacturer": "Server Vendor", "product": "Server", "serial_number": "SYSTEM-123"}
    with pytest.raises(AuditError) as exc:
        parse_fru(raw.replace("Builtin FRU Device (ID 0)", "Board (ID 0)"))
    assert exc.value.code == "PARSE_FAILED" and exc.value.detail == "ambiguous FRU device"
    missing = parse_fru("Product Manufacturer :\nProduct Name : Demo Server\nProduct Serial : DEMO-SN-1")
    assert missing["manufacturer"] == "" and missing["product"] == "Demo Server"


def test_real_ssh_uses_reject_policy_and_fixed_read_only_command(monkeypatch, tmp_path):
    import paramiko

    seen = {}

    class Channel:
        def settimeout(self, value):
            seen["channel_timeout"] = value

        def recv_exit_status(self):
            return 0

    class Stream:
        channel = Channel()

        def __init__(self, payload):
            self.payload = payload

        def read(self, size):
            return self.payload

    class Client:
        def load_host_keys(self, path):
            seen["known_hosts"] = path

        def set_missing_host_key_policy(self, policy):
            seen["policy"] = policy

        def connect(self, **kwargs):
            seen["connect"] = kwargs

        def exec_command(self, command, **kwargs):
            seen["command"] = command
            return None, Stream(b"demo-server-01\n"), Stream(b"")

        def close(self):
            seen["closed"] = True

    monkeypatch.setattr(paramiko, "SSHClient", Client)
    known = tmp_path / "known_hosts"
    assert RealProvider(3, known).linux(Host("demo", "192.0.2.10", "demo", None), "hostname") == "demo-server-01\n"
    assert isinstance(seen["policy"], paramiko.RejectPolicy)
    assert seen["command"] == "LC_ALL=C hostname"
    assert seen["connect"]["look_for_keys"] is False and seen["connect"]["allow_agent"] is False
    assert seen["known_hosts"] == str(known) and seen["closed"]


@pytest.mark.parametrize("error", [OSError("connection refused"), "hostkey"])
def test_real_connection_failure_attempts_once_per_host(monkeypatch, error):
    import paramiko

    attempts = []

    class Client:
        def load_system_host_keys(self):
            pass

        def set_missing_host_key_policy(self, policy):
            assert isinstance(policy, paramiko.RejectPolicy)

        def connect(self, **kwargs):
            attempts.append(kwargs["hostname"])
            if error == "hostkey":
                raise paramiko.SSHException("host key rejected")
            raise error

        def close(self):
            pass

    monkeypatch.setattr(paramiko, "SSHClient", Client)
    hosts = [Host("one", "192.0.2.10", "demo", None),
             Host("two", "192.0.2.11", "demo", None)]
    linux, bmc, failures = collect(hosts, RealProvider(1))
    assert attempts == ["192.0.2.10", "192.0.2.11"]
    assert not bmc and all(row["status"] == "FAILED" for row in linux)
    assert len(failures) == 2 and all(row["field"] == "connection" for row in failures)


def test_wrapped_ssh_banner_timeout_classified_by_cause(monkeypatch):
    import paramiko

    class Client:
        def load_system_host_keys(self):
            pass

        def set_missing_host_key_policy(self, policy):
            pass

        def connect(self, **kwargs):
            raise paramiko.SSHException("banner read failed") from TimeoutError()

        def close(self):
            pass

    monkeypatch.setattr(paramiko, "SSHClient", Client)
    with pytest.raises(AuditError) as exc:
        RealProvider(1).linux(Host("demo", "192.0.2.10", "demo", None), "hostname")
    assert exc.value.code == "SSH_TIMEOUT" and exc.value.field == "connection"


def test_cli_returns_demo_summary(tmp_path, capsys):
    code = main(["--inventory", str(ROOT / "demo" / "hosts.yaml"),
                 "--provider", "mock", "--output", str(tmp_path)])
    assert code == 0
    assert json.loads(capsys.readouterr().out)["hosts"] == 4
